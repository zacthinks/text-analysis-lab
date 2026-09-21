"""Native external-data import operations for Text Analysis Lab (TeAL).

Importers create the first durable artifact in a TeAL lineage.  They are native
core operations rather than Translators because they consume external storage,
not an existing source artifact.

Tabular import uses DuckDB as the parsing/normalization layer and streams Arrow
record batches into TeAL's ordinary ArtifactWriter.  Folder inventory is a
path-only importer: it records filesystem paths and lightweight file attributes
without reading file contents.  Content extraction remains a later Translator
step.
"""

from __future__ import annotations

import json
import os
import re
import tempfile
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path
from typing import TYPE_CHECKING, Any, Literal, get_args

import numpy as np
import pandas as pd

from text_analysis_lab.core.errors import ArtifactError
from text_analysis_lab.core.failure_cleanup import mark_operation_failed_best_effort
from text_analysis_lab.core.ids import next_id
from text_analysis_lab.core.operator import (
    BaseOperator,
    OutputSpec,
    TranslationRequest,
    validate_output_label,
)
from text_analysis_lab.core.types import DEFAULT_OUTPUT_LABEL, StructuralColumn
from text_analysis_lab.core.writer import create_artifact_writer

STRUCTURAL_COLUMNS: frozenset[str] = frozenset(get_args(StructuralColumn))

if TYPE_CHECKING:
    from text_analysis_lab.core.artifact_base import BaseArtifact
    from text_analysis_lab.core.project import Project


TabularFormat = Literal["csv", "jsonl", "parquet"]
ExcelSheetSelector = int | str
ImportDType = str | Mapping[str, str] | None
_MISSING_FIELDS_ERROR = object()
_IMPORT_KINDS = frozenset(
    {
        "read_csv",
        "read_csv_folder",
        "read_excel",
        "read_excel_folder",
        "read_jsonl",
        "read_parquet",
        "folder_inventory",
    }
)
_OPTION_NAME_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")


class ImportOperator(BaseOperator):
    """Frozen snapshot identifying one native import family.

    External paths and parser choices are operation parameters/provenance rather
    than reusable fitted operator state.  The snapshot records only the importer
    family so it remains reconstructible and inspectable after project reopen.
    """

    operation_type = "import"

    def __init__(self, kind: str, *, operator_id: str | None = None) -> None:
        super().__init__(operator_id=operator_id)
        kind = str(kind)
        if kind not in _IMPORT_KINDS:
            raise ValueError(f"Unsupported import kind {kind!r}.")
        self.kind = kind

    def output_specs(
        self,
        *,
        sources: Mapping[str, BaseArtifact],
        request: TranslationRequest,
    ) -> Mapping[str, OutputSpec]:
        if sources:
            raise ArtifactError("Native imports do not consume TeAL source artifacts.")
        label = validate_output_label(
            str(request.params.get("output_label", DEFAULT_OUTPUT_LABEL))
        )
        return {label: OutputSpec(artifact_type="table", lineage_mode="new_key")}

    def to_json_state(self) -> dict[str, Any]:
        return {"kind": self.kind}

    @classmethod
    def from_json_state(cls, state: Mapping[str, Any]) -> ImportOperator:
        return cls(kind=str(state["kind"]))


def read_csv(
    project: Project,
    path: str | Path,
    *,
    text_fields: str | Sequence[str],
    metadata_fields: str | Sequence[str] | None,
    dtype: ImportDType = None,
    batch_size: int = 10_000,
    output_label: str = DEFAULT_OUTPUT_LABEL,
    duckdb_options: Mapping[str, Any] | None = None,
    memo: str | None = None,
) -> BaseArtifact:
    """Import selected CSV fields into a durable TeAL table artifact via DuckDB.

    TeAL always assigns a new 0-based integer ``row_id`` primary key. Only
    ``text_fields`` and ``metadata_fields`` are imported; all other source fields
    are discarded. ``dtype`` may be a pandas dtype string applied to every
    selected field or a mapping from selected field names to dtype strings.
    """
    return _read_tabular(
        project,
        path,
        format="csv",
        text_fields=text_fields,
        metadata_fields=metadata_fields,
        dtype=dtype,
        batch_size=batch_size,
        output_label=output_label,
        duckdb_options=duckdb_options,
        memo=memo,
    )


def read_csv_folder(
    project: Project,
    root: str | Path,
    *,
    text_fields: str | Sequence[str],
    metadata_fields: str | Sequence[str] | None,
    dtype: ImportDType = None,
    pattern: str = "*.csv",
    recursive: bool = True,
    batch_size: int = 10_000,
    output_label: str = DEFAULT_OUTPUT_LABEL,
    duckdb_options: Mapping[str, Any] | None = None,
    memo: str | None = None,
) -> BaseArtifact:
    """Import a deterministic folder of CSV files as one TeAL table artifact.

    Files are discovered beneath ``root`` using ``pattern`` and sorted by POSIX
    relative path before DuckDB reads the explicit file list. TeAL assigns one
    global 0-based ``row_id`` across the combined corpus and attaches exactly two
    generic provenance metadata fields: ``source_file`` (the relative path below
    ``root``) and zero-based ``source_row`` within that CSV. Corpus-specific
    interpretation of filenames belongs in later TeAL transformations. ``dtype``
    may be a pandas dtype string applied to every selected source field or a
    mapping from selected field names to dtype strings.
    """
    root_path = _validate_folder_root(root)
    if not isinstance(pattern, str) or not pattern:
        raise ValueError("pattern must be a non-empty glob string.")
    paths = _inventory_paths(root_path, (pattern,), recursive=recursive)
    if not paths:
        raise ArtifactError(
            f"read_csv_folder matched no files under {str(root_path)!r} "
            f"for pattern {pattern!r}."
        )
    provenance_fields = {"source_file": "source_file", "source_row": "source_row"}
    return _read_csv_files(
        project,
        root_path,
        paths,
        text_fields=text_fields,
        metadata_fields=metadata_fields,
        dtype=dtype,
        provenance_fields=provenance_fields,
        pattern=pattern,
        recursive=recursive,
        batch_size=batch_size,
        output_label=output_label,
        duckdb_options=duckdb_options,
        memo=memo,
    )


def read_excel(
    project: Project,
    path: str | Path,
    *,
    text_fields: str | Sequence[str],
    metadata_fields: str | Sequence[str] | None,
    dtype: ImportDType = None,
    sheets: ExcelSheetSelector | Sequence[ExcelSheetSelector] | None = None,
    header_row: int = 0,
    missing_fields: Any = _MISSING_FIELDS_ERROR,
    batch_size: int = 10_000,
    output_label: str = DEFAULT_OUTPUT_LABEL,
    memo: str | None = None,
) -> BaseArtifact:
    """Import selected rows from one XLSX workbook into a TeAL table artifact.

    ``sheets`` may be a sheet name, zero-based sheet position, a sequence mixing
    names and positions, or ``None`` for every workbook sheet in workbook order.
    Every emitted row receives ``source_sheet``, ``source_sheet_index``, and
    zero-based ``source_row`` metadata. TeAL always assigns one generated global
    ``row_id`` across all selected sheets.

    Requested fields must exist on every selected sheet by default. Supplying
    ``missing_fields`` fills any absent requested field with that scalar value;
    explicitly passing ``None`` therefore fills missing fields with nulls.
    ``dtype`` may be a pandas dtype string applied to every selected field or a
    mapping from selected field names to dtype strings. Excel coercion is applied
    to raw cell values before pandas type inference.
    """
    source_path = _validate_excel_source_file(path)
    return _read_excel_files(
        project,
        root=None,
        paths=[source_path],
        text_fields=text_fields,
        metadata_fields=metadata_fields,
        dtype=dtype,
        sheets=sheets,
        header_row=header_row,
        missing_fields=missing_fields,
        pattern=None,
        recursive=False,
        batch_size=batch_size,
        output_label=output_label,
        memo=memo,
    )


def read_excel_folder(
    project: Project,
    root: str | Path,
    *,
    text_fields: str | Sequence[str],
    metadata_fields: str | Sequence[str] | None,
    dtype: ImportDType = None,
    sheets: ExcelSheetSelector | Sequence[ExcelSheetSelector] | None = None,
    header_row: int = 0,
    missing_fields: Any = _MISSING_FIELDS_ERROR,
    pattern: str = "*.xlsx",
    recursive: bool = True,
    batch_size: int = 10_000,
    output_label: str = DEFAULT_OUTPUT_LABEL,
    memo: str | None = None,
) -> BaseArtifact:
    """Import a deterministic folder of XLSX workbooks as one TeAL artifact.

    Workbooks are traversed in sorted relative-path order. Within each workbook,
    selected sheets are traversed in caller-requested order, or workbook order
    when ``sheets=None``. Generic provenance metadata are ``source_file``,
    ``source_sheet``, ``source_sheet_index``, and zero-based ``source_row``.
    ``dtype`` follows the same scalar-or-mapping coercion contract as
    :func:`read_excel`.
    """
    root_path = _validate_folder_root(root)
    if not isinstance(pattern, str) or not pattern:
        raise ValueError("pattern must be a non-empty glob string.")
    paths = _inventory_paths(root_path, (pattern,), recursive=recursive)
    if not paths:
        raise ArtifactError(
            f"read_excel_folder matched no files under {str(root_path)!r} "
            f"for pattern {pattern!r}."
        )
    for path in paths:
        _validate_excel_source_file(path)
    return _read_excel_files(
        project,
        root=root_path,
        paths=paths,
        text_fields=text_fields,
        metadata_fields=metadata_fields,
        dtype=dtype,
        sheets=sheets,
        header_row=header_row,
        missing_fields=missing_fields,
        pattern=pattern,
        recursive=recursive,
        batch_size=batch_size,
        output_label=output_label,
        memo=memo,
    )


def read_jsonl(
    project: Project,
    path: str | Path,
    *,
    text_fields: str | Sequence[str],
    metadata_fields: str | Sequence[str] | None,
    dtype: ImportDType = None,
    batch_size: int = 10_000,
    output_label: str = DEFAULT_OUTPUT_LABEL,
    duckdb_options: Mapping[str, Any] | None = None,
    memo: str | None = None,
) -> BaseArtifact:
    """Import selected JSONL fields into a durable TeAL table artifact via DuckDB.

    TeAL always assigns a new 0-based integer ``row_id`` primary key. Only
    ``text_fields`` and ``metadata_fields`` are imported; all other source fields
    are discarded. ``dtype`` may be a pandas dtype string applied to every
    selected field or a mapping from selected field names to dtype strings.
    """
    options = dict(duckdb_options or {})
    if "format" in options:
        raise ValueError(
            "read_jsonl fixes DuckDB format='newline_delimited'; do not pass a format option."
        )
    return _read_tabular(
        project,
        path,
        format="jsonl",
        text_fields=text_fields,
        metadata_fields=metadata_fields,
        dtype=dtype,
        batch_size=batch_size,
        output_label=output_label,
        duckdb_options=options,
        memo=memo,
    )


def read_parquet(
    project: Project,
    path: str | Path,
    *,
    text_fields: str | Sequence[str],
    metadata_fields: str | Sequence[str] | None,
    dtype: ImportDType = None,
    batch_size: int = 10_000,
    output_label: str = DEFAULT_OUTPUT_LABEL,
    duckdb_options: Mapping[str, Any] | None = None,
    memo: str | None = None,
) -> BaseArtifact:
    """Import selected Parquet fields into a durable TeAL table artifact via DuckDB.

    TeAL always assigns a new 0-based integer ``row_id`` primary key. Only
    ``text_fields`` and ``metadata_fields`` are imported; all other source fields
    are discarded. ``dtype`` may be a pandas dtype string applied to every
    selected field or a mapping from selected field names to dtype strings.
    """
    return _read_tabular(
        project,
        path,
        format="parquet",
        text_fields=text_fields,
        metadata_fields=metadata_fields,
        dtype=dtype,
        batch_size=batch_size,
        output_label=output_label,
        duckdb_options=duckdb_options,
        memo=memo,
    )


def folder_inventory(
    project: Project,
    root: str | Path,
    *,
    patterns: str | Sequence[str] = "*",
    recursive: bool = True,
    batch_size: int = 10_000,
    output_label: str = DEFAULT_OUTPUT_LABEL,
    memo: str | None = None,
) -> BaseArtifact:
    """Inventory matching files as a path-only TeAL table artifact.

    The output primary key is generated ``file_id`` in deterministic sorted
    relative-path order.  Matching the same file through multiple patterns does
    not duplicate it.  File contents are intentionally not read here.
    """
    root_path = _validate_folder_root(root)
    pattern_values = _validate_patterns(patterns)
    batch_size = _validate_batch_size(batch_size)
    output_label = validate_output_label(output_label)
    paths = _inventory_paths(root_path, pattern_values, recursive=recursive)
    if not paths:
        raise ArtifactError(
            f"folder_inventory matched no files under {str(root_path)!r} "
            f"for patterns {list(pattern_values)!r}."
        )

    request = {
        "root": str(root_path),
        "patterns": list(pattern_values),
        "recursive": bool(recursive),
        "batch_size": batch_size,
        "output_label": output_label,
    }
    external_source = {
        "kind": "folder",
        "path": str(root_path),
        "matched_file_count": len(paths),
    }

    def payloads() -> Iterable[Mapping[str, Any]]:
        for start in range(0, len(paths), batch_size):
            batch_paths = paths[start : start + batch_size]
            keys = pd.DataFrame(
                {"file_id": np.arange(start, start + len(batch_paths), dtype="int64")}
            )
            rows: list[dict[str, Any]] = []
            for path in batch_paths:
                stat = path.stat()
                relative = path.relative_to(root_path).as_posix()
                rows.append(
                    {
                        "path": str(path),
                        "relative_path": relative,
                        "file_name": path.name,
                        "stem": path.stem,
                        "extension": path.suffix,
                        "size_bytes": int(stat.st_size),
                    }
                )
            yield {"keys": keys, "data": pd.DataFrame(rows)}

    return _execute_import(
        project,
        kind="folder_inventory",
        output_label=output_label,
        request=request,
        external_source=external_source,
        payloads=payloads(),
        memo=memo,
    )


def _read_csv_files(
    project: Project,
    root: Path,
    paths: Sequence[Path],
    *,
    text_fields: str | Sequence[str],
    metadata_fields: str | Sequence[str] | None,
    dtype: ImportDType,
    provenance_fields: Mapping[str, str],
    pattern: str,
    recursive: bool,
    batch_size: int,
    output_label: str,
    duckdb_options: Mapping[str, Any] | None,
    memo: str | None,
) -> BaseArtifact:
    batch_size = _validate_batch_size(batch_size)
    output_label = validate_output_label(output_label)
    options = _validate_duckdb_options(duckdb_options)
    if "filename" in options:
        raise ValueError(
            "read_csv_folder manages DuckDB filename provenance internally; "
            "do not pass a filename option."
        )

    duckdb = _import_duckdb()
    con = duckdb.connect(database=":memory:")
    try:
        con.execute("SET preserve_insertion_order = true")
        path_values = [path.as_posix() for path in paths]
        source_sql = _duckdb_source_sql_many(
            "csv", path_values, {**options, "filename": True}
        )
        source_relation = con.sql(f"SELECT * FROM {source_sql}")
        columns = tuple(str(column) for column in source_relation.columns)
        duckdb_types = tuple(str(dtype) for dtype in source_relation.dtypes)
        if "filename" not in columns:
            raise ArtifactError(
                "DuckDB CSV folder import did not expose filename provenance."
            )
        source_columns = tuple(column for column in columns if column != "filename")
        source_types = tuple(
            dtype
            for column, dtype in zip(columns, duckdb_types, strict=True)
            if column != "filename"
        )
        plan = _plan_tabular_columns(
            source_columns,
            text_fields=text_fields,
            metadata_fields=metadata_fields,
        )
        dtype_map = _normalize_import_dtype(
            dtype,
            selected_fields=(*plan["text_fields"], *plan["metadata_fields"]),
        )
        collisions = sorted(
            set(source_columns).intersection(provenance_fields.values())
        )
        if collisions:
            raise ArtifactError(
                "CSV-folder provenance field name(s) collide with selected source columns: "
                f"{collisions}. Rename/remove the colliding source column before import."
            )

        selected_fields = (*plan["text_fields"], *plan["metadata_fields"], "filename")
        projection = ", ".join(
            _quote_duckdb_identifier(field) for field in selected_fields
        )
        relation = con.sql(f"SELECT {projection} FROM {source_sql}")

        request = {
            "root": str(root),
            "pattern": pattern,
            "recursive": bool(recursive),
            "matched_files": [path.relative_to(root).as_posix() for path in paths],
            "format": "csv",
            "primary_key": ["row_id"],
            "generated_primary_key": True,
            "text_fields": list(plan["text_fields"]),
            "metadata_fields": [*plan["metadata_fields"], *provenance_fields.values()],
            "source_metadata_fields": list(plan["metadata_fields"]),
            "provenance_fields": dict(provenance_fields),
            "discarded_fields": list(plan["discarded_fields"]),
            "batch_size": batch_size,
            "output_label": output_label,
            "duckdb_options": dict(options),
            "reader_engine": {"name": "duckdb", "version": str(duckdb.__version__)},
            "detected_schema": [
                {"name": name, "duckdb_type": dtype}
                for name, dtype in zip(source_columns, source_types, strict=True)
            ],
        }
        if dtype_map:
            request["dtype"] = dict(dtype_map)
        external_source = {
            "kind": "folder",
            "path": str(root),
            "format": "csv",
            "matched_file_count": len(paths),
            "matched_files": [path.relative_to(root).as_posix() for path in paths],
        }

        reader = relation.to_arrow_reader(batch_size=batch_size)
        per_file_rows: dict[str, int] = {}

        def payloads() -> Iterable[Mapping[str, Any]]:
            next_row_id = 0
            for record_batch in reader:
                frame = record_batch.to_pandas().reset_index(drop=True)
                if frame.empty:
                    continue
                frame = _coerce_import_dtypes(
                    frame,
                    dtype_map,
                    context="CSV folder import",
                )
                keys = pd.DataFrame(
                    {
                        "row_id": np.arange(
                            next_row_id,
                            next_row_id + len(frame),
                            dtype="int64",
                        )
                    }
                )
                next_row_id += len(frame)
                metadata = frame.loc[:, list(plan["metadata_fields"])].copy()
                source_files: list[str] = []
                source_rows: list[int] = []
                for raw_filename in frame["filename"].astype(str):
                    source_path = Path(raw_filename).resolve()
                    try:
                        relative = source_path.relative_to(root).as_posix()
                    except ValueError as exc:
                        raise ArtifactError(
                            f"DuckDB returned source path outside CSV-folder root: {source_path}."
                        ) from exc
                    key = source_path.as_posix()
                    source_row = per_file_rows.get(key, 0)
                    per_file_rows[key] = source_row + 1
                    source_files.append(relative)
                    source_rows.append(source_row)
                metadata["source_file"] = source_files
                metadata["source_row"] = np.asarray(source_rows, dtype="int64")

                yield {
                    "keys": keys,
                    "data": frame.loc[:, list(plan["text_fields"])].copy(),
                    "metadata": metadata,
                }

        return _execute_import(
            project,
            kind="read_csv_folder",
            output_label=output_label,
            request=request,
            external_source=external_source,
            payloads=payloads(),
            memo=memo,
        )
    finally:
        con.close()


def _read_excel_files(
    project: Project,
    *,
    root: Path | None,
    paths: Sequence[Path],
    text_fields: str | Sequence[str],
    metadata_fields: str | Sequence[str] | None,
    dtype: ImportDType,
    sheets: ExcelSheetSelector | Sequence[ExcelSheetSelector] | None,
    header_row: int,
    missing_fields: Any,
    pattern: str | None,
    recursive: bool,
    batch_size: int,
    output_label: str,
    memo: str | None,
) -> BaseArtifact:
    batch_size = _validate_batch_size(batch_size)
    output_label = validate_output_label(output_label)
    header_row = _validate_header_row(header_row)
    sheet_selectors = _normalize_sheet_selectors(sheets)
    missing_policy = _validate_missing_fields_policy(missing_fields)

    text = _normalize_column_names(text_fields, name="text_fields")
    if not text:
        raise ValueError("text_fields must contain at least one source column.")
    metadata = _normalize_column_names(metadata_fields, name="metadata_fields")
    _validate_selected_field_names(text, metadata)
    dtype_map = _normalize_import_dtype(
        dtype,
        selected_fields=(*text, *metadata),
    )

    provenance_fields = ["source_sheet", "source_sheet_index", "source_row"]
    if root is not None:
        provenance_fields.insert(0, "source_file")
    collisions = sorted({*text, *metadata}.intersection(provenance_fields))
    if collisions:
        raise ArtifactError(
            "Excel provenance field name(s) collide with requested source fields: "
            f"{collisions}. Rename/remove the colliding source field before import."
        )

    openpyxl = _import_openpyxl()
    inspections: list[dict[str, Any]] = []
    workbook_plans: list[dict[str, Any]] = []
    for path in paths:
        workbook = openpyxl.load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
        )
        try:
            resolved_sheets = _resolve_excel_sheets(
                workbook.sheetnames,
                sheet_selectors,
                source_path=path,
            )
            sheet_plans: list[dict[str, Any]] = []
            for sheet_index, sheet_name in resolved_sheets:
                worksheet = workbook[sheet_name]
                columns = _excel_header_columns(worksheet, header_row=header_row)
                available = tuple(columns)
                requested = (*text, *metadata)
                missing = tuple(
                    field for field in requested if field not in set(available)
                )
                if missing and missing_policy["mode"] == "error":
                    relative = (
                        path.relative_to(root).as_posix()
                        if root is not None
                        else str(path)
                    )
                    raise ArtifactError(
                        f"Excel source {relative!r}, sheet {sheet_name!r} is missing "
                        f"requested field(s) {list(missing)}; available columns are "
                        f"{list(available)}. Pass missing_fields=None to fill missing "
                        "requested fields with nulls, or pass another scalar default value."
                    )
                discarded = tuple(
                    field for field in available if field not in set(requested)
                )
                plan = {
                    "sheet_name": sheet_name,
                    "sheet_index": int(sheet_index),
                    "columns": available,
                    "missing_fields": missing,
                    "discarded_fields": discarded,
                }
                sheet_plans.append(plan)
                inspections.append(
                    {
                        **(
                            {"source_file": path.relative_to(root).as_posix()}
                            if root is not None
                            else {}
                        ),
                        "sheet": sheet_name,
                        "sheet_index": int(sheet_index),
                        "columns": list(available),
                        "missing_fields": list(missing),
                        "discarded_fields": list(discarded),
                    }
                )
            workbook_plans.append({"path": path, "sheets": sheet_plans})
        finally:
            workbook.close()

    requested_sheets: list[int | str] | None
    if sheet_selectors is None:
        requested_sheets = None
    else:
        requested_sheets = list(sheet_selectors)

    common_request: dict[str, Any] = {
        "format": "xlsx",
        "primary_key": ["row_id"],
        "generated_primary_key": True,
        "text_fields": list(text),
        "metadata_fields": [*metadata, *provenance_fields],
        "source_metadata_fields": list(metadata),
        "provenance_fields": provenance_fields,
        "sheets": requested_sheets,
        "header_row": header_row,
        "missing_fields": missing_policy,
        "batch_size": batch_size,
        "output_label": output_label,
        "reader_engine": {
            "name": "openpyxl",
            "version": str(openpyxl.__version__),
            "read_only": True,
            "data_only": True,
        },
        "formula_policy": "cached_values",
        "detected_sheets": inspections,
    }
    if dtype_map:
        common_request["dtype"] = dict(dtype_map)
    if root is None:
        source_path = paths[0]
        request = {"path": str(source_path), **common_request}
        stat = source_path.stat()
        external_source = {
            "kind": "file",
            "path": str(source_path),
            "format": "xlsx",
            "size_bytes": int(stat.st_size),
            "modified_ns": int(stat.st_mtime_ns),
        }
        kind = "read_excel"
    else:
        matched_files = [path.relative_to(root).as_posix() for path in paths]
        request = {
            "root": str(root),
            "pattern": pattern,
            "recursive": bool(recursive),
            "matched_files": matched_files,
            **common_request,
        }
        external_source = {
            "kind": "folder",
            "path": str(root),
            "format": "xlsx",
            "matched_file_count": len(paths),
            "matched_files": matched_files,
        }
        kind = "read_excel_folder"

    def payloads() -> Iterable[Mapping[str, Any]]:
        next_row_id = 0
        for workbook_plan in workbook_plans:
            path = workbook_plan["path"]
            workbook = openpyxl.load_workbook(
                filename=path,
                read_only=True,
                data_only=True,
            )
            try:
                for sheet_plan in workbook_plan["sheets"]:
                    sheet_name = str(sheet_plan["sheet_name"])
                    sheet_index = int(sheet_plan["sheet_index"])
                    columns = tuple(str(column) for column in sheet_plan["columns"])
                    column_positions = {
                        column: index for index, column in enumerate(columns)
                    }
                    worksheet = workbook[sheet_name]
                    row_iter = worksheet.iter_rows(
                        min_row=header_row + 2,
                        values_only=True,
                    )
                    source_row = 0
                    batch_rows: list[dict[str, Any]] = []
                    batch_source_rows: list[int] = []

                    def flush(
                        *,
                        path: Path = path,
                        sheet_name: str = sheet_name,
                        sheet_index: int = sheet_index,
                    ) -> Mapping[str, Any] | None:
                        nonlocal next_row_id, batch_rows, batch_source_rows
                        if not batch_rows:
                            return None
                        frame = _frame_from_import_rows(
                            batch_rows,
                            columns=[*text, *metadata],
                            dtype_map=dtype_map,
                            context=(
                                f"Excel source {str(path)!r}, sheet {sheet_name!r}"
                            ),
                        )
                        count = len(frame)
                        keys = pd.DataFrame(
                            {
                                "row_id": np.arange(
                                    next_row_id,
                                    next_row_id + count,
                                    dtype="int64",
                                )
                            }
                        )
                        next_row_id += count
                        out_metadata = frame.loc[:, list(metadata)].copy()
                        if root is not None:
                            out_metadata["source_file"] = path.relative_to(
                                root
                            ).as_posix()
                        out_metadata["source_sheet"] = sheet_name
                        out_metadata["source_sheet_index"] = np.full(
                            count, sheet_index, dtype="int64"
                        )
                        out_metadata["source_row"] = np.asarray(
                            batch_source_rows, dtype="int64"
                        )
                        payload: dict[str, Any] = {
                            "keys": keys,
                            "data": frame.loc[:, list(text)].copy(),
                            "metadata": out_metadata,
                        }
                        batch_rows = []
                        batch_source_rows = []
                        return payload

                    for raw_values in row_iter:
                        values = tuple(raw_values)
                        if _excel_row_is_blank(values, width=len(columns)):
                            continue
                        row: dict[str, Any] = {}
                        for field in (*text, *metadata):
                            position = column_positions.get(field)
                            if position is None:
                                row[field] = missing_policy["value"]
                            elif position < len(values):
                                row[field] = values[position]
                            else:
                                row[field] = None
                        batch_rows.append(row)
                        batch_source_rows.append(source_row)
                        source_row += 1
                        if len(batch_rows) >= batch_size:
                            payload = flush()
                            if payload is not None:
                                yield payload
                    payload = flush()
                    if payload is not None:
                        yield payload
            finally:
                workbook.close()

    return _execute_import(
        project,
        kind=kind,
        output_label=output_label,
        request=request,
        external_source=external_source,
        payloads=payloads(),
        memo=memo,
    )


def _normalize_sheet_selectors(
    sheets: ExcelSheetSelector | Sequence[ExcelSheetSelector] | None,
) -> tuple[ExcelSheetSelector, ...] | None:
    if sheets is None:
        return None
    if isinstance(sheets, bool):
        raise TypeError(
            "sheets entries must be sheet names or zero-based integer positions."
        )
    if isinstance(sheets, (str, int)):
        values: tuple[Any, ...] = (sheets,)
    elif isinstance(sheets, Sequence):
        values = tuple(sheets)
    else:
        raise TypeError(
            "sheets must be a sheet name, integer position, sequence of names/positions, or None."
        )
    if not values:
        raise ValueError(
            "sheets cannot be an empty sequence; use None to import all sheets."
        )
    normalized: list[ExcelSheetSelector] = []
    for value in values:
        if isinstance(value, bool) or not isinstance(value, (str, int)):
            raise TypeError(
                "sheets entries must be sheet names or zero-based integer positions."
            )
        if isinstance(value, int) and value < 0:
            raise ValueError(
                "Excel sheet positions must be zero-based non-negative integers."
            )
        if isinstance(value, str) and not value:
            raise ValueError("Excel sheet names cannot be empty.")
        normalized.append(value)
    return tuple(normalized)


def _resolve_excel_sheets(
    sheet_names: Sequence[str],
    selectors: tuple[ExcelSheetSelector, ...] | None,
    *,
    source_path: Path,
) -> tuple[tuple[int, str], ...]:
    names = tuple(str(name) for name in sheet_names)
    if not names:
        raise ArtifactError(f"Excel workbook exposes no sheets: {source_path}.")
    if selectors is None:
        return tuple(enumerate(names))

    resolved: list[tuple[int, str]] = []
    seen_indices: set[int] = set()
    for selector in selectors:
        if isinstance(selector, int):
            if selector >= len(names):
                raise ArtifactError(
                    f"Excel workbook {str(source_path)!r} has {len(names)} sheet(s); "
                    f"requested sheet position {selector} does not exist."
                )
            index = selector
            name = names[index]
        else:
            if selector not in names:
                raise ArtifactError(
                    f"Excel workbook {str(source_path)!r} has no sheet {selector!r}; "
                    f"available sheets are {list(names)}."
                )
            index = names.index(selector)
            name = names[index]
        if index in seen_indices:
            raise ArtifactError(
                f"Excel workbook {str(source_path)!r} resolves the sheets selection "
                f"to duplicate sheet {name!r} (position {index})."
            )
        seen_indices.add(index)
        resolved.append((index, name))
    return tuple(resolved)


def _excel_header_columns(worksheet: Any, *, header_row: int) -> tuple[str, ...]:
    iterator = worksheet.iter_rows(
        min_row=header_row + 1,
        max_row=header_row + 1,
        values_only=True,
    )
    raw = tuple(next(iterator, ()))
    while raw and raw[-1] is None:
        raw = raw[:-1]
    if not raw:
        return ()
    columns: list[str] = []
    for index, value in enumerate(raw):
        if value is None or str(value) == "":
            columns.append(f"Unnamed: {index}")
        else:
            columns.append(str(value))
    if len(set(columns)) != len(columns):
        duplicates = sorted({name for name in columns if columns.count(name) > 1})
        raise ArtifactError(
            f"Excel sheet contains duplicate header names after normalization: {duplicates}."
        )
    return tuple(columns)


def _excel_row_is_blank(values: Sequence[Any], *, width: int) -> bool:
    if width <= 0:
        return all(value is None for value in values)
    considered = values[:width]
    if len(considered) < width:
        considered = (*considered, *((None,) * (width - len(considered))))
    return all(value is None for value in considered)


def _validate_selected_field_names(
    text_fields: Sequence[str],
    metadata_fields: Sequence[str],
) -> None:
    overlap = sorted(set(text_fields).intersection(metadata_fields))
    if overlap:
        raise ArtifactError(
            f"Source columns cannot be both text_fields and metadata_fields: {overlap}."
        )
    selected = set(text_fields).union(metadata_fields)
    reserved = sorted(selected.intersection(STRUCTURAL_COLUMNS))
    if reserved:
        raise ArtifactError(
            f"Selected source field(s) use TeAL-reserved structural names: {reserved}."
        )
    if "row_id" in selected:
        raise ArtifactError(
            "Source field 'row_id' cannot be imported because TeAL reserves row_id "
            "for the generated primary key of tabular imports. Rename that source "
            "field before import or leave it unselected."
        )


def _validate_header_row(value: int) -> int:
    if isinstance(value, bool) or not isinstance(value, int) or value < 0:
        raise ValueError("header_row must be a zero-based non-negative integer.")
    return int(value)


def _validate_missing_fields_policy(value: Any) -> dict[str, Any]:
    if value is _MISSING_FIELDS_ERROR:
        return {"mode": "error"}
    if isinstance(value, np.generic):
        value = value.item()
    if value is not None and not isinstance(value, (str, bool, int, float)):
        raise TypeError(
            "missing_fields must be omitted to raise on missing requested fields, "
            "or be a scalar string/boolean/number/None used as the fill value."
        )
    if isinstance(value, float) and not np.isfinite(value):
        raise ValueError("missing_fields numeric fill values must be finite.")
    return {"mode": "fill", "value": value}


def _validate_excel_source_file(path: str | Path) -> Path:
    resolved = _validate_source_file(path)
    if resolved.suffix.lower() != ".xlsx":
        raise ArtifactError(
            f"Excel import currently supports .xlsx workbooks only: {resolved}."
        )
    return resolved


def _import_openpyxl():
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - required package in real env
        raise RuntimeError(
            "openpyxl is required for TeAL Excel imports. Install project dependencies."
        ) from exc
    return openpyxl


def _read_tabular(
    project: Project,
    path: str | Path,
    *,
    format: TabularFormat,
    text_fields: str | Sequence[str],
    metadata_fields: str | Sequence[str] | None,
    dtype: ImportDType,
    batch_size: int,
    output_label: str,
    duckdb_options: Mapping[str, Any] | None,
    memo: str | None,
) -> BaseArtifact:
    source_path = _validate_source_file(path)
    batch_size = _validate_batch_size(batch_size)
    output_label = validate_output_label(output_label)
    options = _validate_duckdb_options(duckdb_options)
    source_sql = _duckdb_source_sql(format, source_path, options)

    duckdb = _import_duckdb()
    con = duckdb.connect(database=":memory:")
    try:
        con.execute("SET preserve_insertion_order = true")
        source_relation = con.sql(f"SELECT * FROM {source_sql}")
        columns = tuple(str(column) for column in source_relation.columns)
        duckdb_types = tuple(str(dtype) for dtype in source_relation.dtypes)
        plan = _plan_tabular_columns(
            columns,
            text_fields=text_fields,
            metadata_fields=metadata_fields,
        )
        dtype_map = _normalize_import_dtype(
            dtype,
            selected_fields=(*plan["text_fields"], *plan["metadata_fields"]),
        )

        selected_fields = (*plan["text_fields"], *plan["metadata_fields"])
        projection = ", ".join(
            _quote_duckdb_identifier(field) for field in selected_fields
        )
        relation = con.sql(f"SELECT {projection} FROM {source_sql}")

        request = {
            "path": str(source_path),
            "format": format,
            "primary_key": ["row_id"],
            "generated_primary_key": True,
            "text_fields": list(plan["text_fields"]),
            "metadata_fields": list(plan["metadata_fields"]),
            "discarded_fields": list(plan["discarded_fields"]),
            "batch_size": batch_size,
            "output_label": output_label,
            "duckdb_options": dict(options),
            "reader_engine": {"name": "duckdb", "version": str(duckdb.__version__)},
            "detected_schema": [
                {"name": name, "duckdb_type": dtype}
                for name, dtype in zip(columns, duckdb_types, strict=True)
            ],
        }
        if dtype_map:
            request["dtype"] = dict(dtype_map)
        stat = source_path.stat()
        external_source = {
            "kind": "file",
            "path": str(source_path),
            "format": format,
            "size_bytes": int(stat.st_size),
            "modified_ns": int(stat.st_mtime_ns),
        }

        reader = relation.to_arrow_reader(batch_size=batch_size)

        def payloads() -> Iterable[Mapping[str, Any]]:
            next_row_id = 0
            for record_batch in reader:
                frame = record_batch.to_pandas().reset_index(drop=True)
                if frame.empty:
                    continue
                frame = _coerce_import_dtypes(
                    frame,
                    dtype_map,
                    context=f"{format.upper()} import {str(source_path)!r}",
                )
                keys = pd.DataFrame(
                    {
                        "row_id": np.arange(
                            next_row_id,
                            next_row_id + len(frame),
                            dtype="int64",
                        )
                    }
                )
                next_row_id += len(frame)

                payload: dict[str, Any] = {
                    "keys": keys,
                    "data": frame.loc[:, list(plan["text_fields"])].copy(),
                }
                if plan["metadata_fields"]:
                    payload["metadata"] = frame.loc[
                        :, list(plan["metadata_fields"])
                    ].copy()
                yield payload

        return _execute_import(
            project,
            kind=f"read_{format}",
            output_label=output_label,
            request=request,
            external_source=external_source,
            payloads=payloads(),
            memo=memo,
        )
    finally:
        con.close()


def _execute_import(
    project: Project,
    *,
    kind: str,
    output_label: str,
    request: Mapping[str, Any],
    external_source: Mapping[str, Any],
    payloads: Iterable[Mapping[str, Any]],
    memo: str | None,
) -> BaseArtifact:
    operator = ImportOperator(kind)
    operator_id = next_id(project.storage.manifest_path, "operator")
    operator.assign_operator_id(operator_id)
    project.catalog.register_operator(
        operator_id=operator_id,
        operation_type="import",
        snapshot_status="pending",
    )
    try:
        operator.save_to_dir(
            project.storage.operator_dir(operator_id),
            operator_id=operator_id,
        )
        project.catalog.mark_operator_serialized(operator_id)
    except Exception:
        project.catalog.mark_operator_snapshot_failed(operator_id)
        raise

    operation_id = next_id(project.storage.manifest_path, "operation")
    operation_dir = project.storage.operation_dir(operation_id)
    operation_dir.mkdir(parents=True, exist_ok=True)
    project.catalog.register_operation(
        operation_id=operation_id,
        operation_type="import",
        operator_id=operator_id,
        status="incomplete",
    )

    artifact_id = next_id(project.storage.manifest_path, "artifact")
    project.catalog.register_artifact(
        artifact_id=artifact_id,
        artifact_type="table",
        label=output_label,
        lineage_mode="new_key",
        status="incomplete",
        basis_artifact_ids=(),
    )
    project.catalog.add_operation_output(
        operation_id,
        output_label,
        artifact_id,
        ordinal=0,
    )

    writer = create_artifact_writer(
        artifact_type="table",
        artifact_dir=project.storage.artifact_dir(artifact_id),
        artifact_id=artifact_id,
        label=output_label,
        operation_id=operation_id,
        lineage_mode="new_key",
        basis_artifact_ids=(),
    )

    descriptor: dict[str, Any] = {
        "schema_version": 1,
        "operation_id": operation_id,
        "operation_type": "import",
        "operator_id": operator_id,
        "import_kind": kind,
        "status": "incomplete",
        "resumable": False,
        "sources": {},
        "external_source": dict(external_source),
        "output_artifact_ids": {output_label: artifact_id},
        "request": dict(request),
    }
    _write_json(operation_dir / "operation.json", descriptor)

    try:
        if memo is not None:
            project.catalog.add_memo(
                target_type="operation",
                target_id=operation_id,
                body=memo,
            )

        wrote_any = False
        for payload in payloads:
            keys = payload.get("keys")
            if isinstance(keys, pd.DataFrame) and keys.empty:
                continue
            writer.write(payload)
            wrote_any = True

        if not wrote_any:
            raise ArtifactError("Import source contains no rows to materialize.")

        writer.finalize()
        project.catalog.mark_artifact_complete(artifact_id)
        project.catalog.mark_operation_complete(operation_id)
        descriptor["status"] = "complete"
        _write_json(operation_dir / "operation.json", descriptor)
        project.storage.touch_manifest()
        return project.get_artifact(artifact_id)
    except Exception as exc:
        mark_operation_failed_best_effort(
            project,
            writer=writer,
            artifact_id=artifact_id,
            operation_id=operation_id,
            error=exc,
        )
        descriptor["status"] = "failed"
        descriptor["error"] = f"{exc.__class__.__name__}: {exc}"
        _write_json(operation_dir / "operation.json", descriptor)
        raise


def _plan_tabular_columns(
    columns: Sequence[str],
    *,
    text_fields: str | Sequence[str],
    metadata_fields: str | Sequence[str] | None,
) -> dict[str, Any]:
    source_columns = tuple(str(column) for column in columns)
    if not source_columns:
        raise ArtifactError("Tabular import source exposes no columns.")
    if len(set(source_columns)) != len(source_columns):
        raise ArtifactError("Tabular import source contains duplicate column names.")
    text = _normalize_column_names(text_fields, name="text_fields")
    if not text:
        raise ValueError("text_fields must contain at least one source column.")
    metadata = _normalize_column_names(metadata_fields, name="metadata_fields")
    _require_known_columns(text, source_columns, name="text_fields")
    _require_known_columns(metadata, source_columns, name="metadata_fields")

    overlap = sorted(set(text).intersection(metadata))
    if overlap:
        raise ArtifactError(
            f"Source columns cannot be both text_fields and metadata_fields: {overlap}."
        )

    selected = set(text).union(metadata)
    reserved = sorted(selected.intersection(STRUCTURAL_COLUMNS))
    if reserved:
        raise ArtifactError(
            f"Selected source field(s) use TeAL-reserved structural names: {reserved}."
        )
    if "row_id" in selected:
        raise ArtifactError(
            "Source field 'row_id' cannot be imported because TeAL reserves row_id "
            "for the generated primary key of tabular imports. Rename that source "
            "field before import or leave it unselected."
        )

    discarded = tuple(column for column in source_columns if column not in selected)
    return {
        "primary_key": ("row_id",),
        "text_fields": tuple(text),
        "metadata_fields": tuple(metadata),
        "discarded_fields": discarded,
    }


def _normalize_import_dtype(
    dtype: ImportDType,
    *,
    selected_fields: Sequence[str],
) -> dict[str, str]:
    """Normalize an optional import-time dtype coercion specification.

    ``None`` preserves the source reader's normal inferred/native types. A scalar
    dtype string applies to every selected source field, while a mapping applies
    only to the named selected fields. Dtype names follow pandas ``astype``
    conventions so nullable dtypes such as ``string``, ``Int64``, ``Float64``,
    and ``boolean`` are supported.

    The normalized result is deliberately a plain string mapping so it can be
    recorded directly in operation provenance.
    """
    selected = tuple(str(field) for field in selected_fields)
    selected_set = set(selected)

    if dtype is None:
        return {}

    if isinstance(dtype, str):
        target = _validate_import_dtype_name(dtype)
        return {field: target for field in selected}

    if not isinstance(dtype, Mapping):
        raise TypeError(
            "dtype must be a dtype string, a mapping of selected field names to "
            "dtype strings, or None."
        )

    normalized: dict[str, str] = {}
    for raw_field, raw_target in dtype.items():
        if not isinstance(raw_field, str) or not raw_field:
            raise TypeError("dtype mapping keys must be non-empty field-name strings.")
        if raw_field not in selected_set:
            raise ArtifactError(
                f"dtype specifies unselected field {raw_field!r}; selected source "
                f"fields are {list(selected)}."
            )
        if not isinstance(raw_target, str):
            raise TypeError(
                f"dtype for field {raw_field!r} must be a dtype string, got "
                f"{type(raw_target).__name__}."
            )
        normalized[raw_field] = _validate_import_dtype_name(raw_target)
    return normalized


def _validate_import_dtype_name(value: str) -> str:
    target = str(value).strip()
    if not target:
        raise ValueError("dtype strings cannot be empty.")
    try:
        pd.api.types.pandas_dtype(target)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Unsupported import dtype {target!r}.") from exc
    return target


def _coerce_import_dtypes(
    frame: pd.DataFrame,
    dtype_map: Mapping[str, str],
    *,
    context: str,
) -> pd.DataFrame:
    """Apply strict import-time dtype coercions before artifact materialization."""
    if not dtype_map:
        return frame

    result = frame.copy()
    for field, target in dtype_map.items():
        if field not in result.columns:
            raise ArtifactError(
                f"Import dtype coercion expected field {field!r} in {context}, but "
                f"available columns are {list(result.columns)}."
            )
        try:
            result[field] = result[field].astype(target)
        except (TypeError, ValueError, OverflowError) as exc:
            raise ArtifactError(
                f"Could not coerce imported field {field!r} to dtype {target!r} "
                f"in {context}: {exc}"
            ) from exc
    return result


def _frame_from_import_rows(
    rows: Sequence[Mapping[str, Any]],
    *,
    columns: Sequence[str],
    dtype_map: Mapping[str, str],
    context: str,
) -> pd.DataFrame:
    """Build a frame while applying dtypes to raw values before pandas inference.

    Excel cells arrive as Python objects (for example ``datetime.time`` or
    integers). Constructing an untyped DataFrame first can coerce those values
    (for example ``7`` plus a missing value becomes ``7.0``), which loses the
    source representation before a later string cast. This helper assigns typed
    arrays from the original cell values instead.
    """
    frame = pd.DataFrame(rows, columns=list(columns))
    if not dtype_map:
        return frame

    for field, target in dtype_map.items():
        try:
            frame[field] = pd.Series(
                [row.get(field) for row in rows],
                dtype=target,
            )
        except (TypeError, ValueError, OverflowError) as exc:
            raise ArtifactError(
                f"Could not coerce imported field {field!r} to dtype {target!r} "
                f"in {context}: {exc}"
            ) from exc
    return frame


def _quote_duckdb_identifier(value: str) -> str:
    return '"' + str(value).replace('"', '""') + '"'


def _normalize_column_names(
    value: str | Sequence[str] | None,
    *,
    name: str,
) -> tuple[str, ...]:
    if value is None:
        return ()
    if isinstance(value, str):
        values = (value,)
    elif isinstance(value, Sequence):
        values = tuple(str(item) for item in value)
    else:
        raise TypeError(f"{name} must be a string, sequence of strings, or None.")
    if any(not item for item in values):
        raise ValueError(f"{name} cannot contain empty column names.")
    if len(set(values)) != len(values):
        raise ValueError(f"{name} cannot contain duplicate column names.")
    return values


def _require_known_columns(
    requested: Sequence[str],
    available: Sequence[str],
    *,
    name: str,
) -> None:
    available_set = set(available)
    missing = [column for column in requested if column not in available_set]
    if missing:
        raise ArtifactError(
            f"Unknown {name} column(s) {missing}; available columns are {list(available)}."
        )


def _validate_source_file(path: str | Path) -> Path:
    resolved = Path(path).expanduser().resolve()
    if not resolved.exists():
        raise FileNotFoundError(resolved)
    if not resolved.is_file():
        raise ArtifactError(f"Import source is not a file: {resolved}.")
    return resolved


def _validate_folder_root(path: str | Path) -> Path:
    resolved = Path(path).expanduser().resolve()
    if not resolved.exists():
        raise FileNotFoundError(resolved)
    if not resolved.is_dir():
        raise ArtifactError(f"Folder inventory root is not a directory: {resolved}.")
    return resolved


def _validate_patterns(patterns: str | Sequence[str]) -> tuple[str, ...]:
    if isinstance(patterns, str):
        values = (patterns,)
    elif isinstance(patterns, Sequence):
        values = tuple(str(item) for item in patterns)
    else:
        raise TypeError("patterns must be a glob string or sequence of glob strings.")
    if not values or any(not value for value in values):
        raise ValueError("patterns must contain at least one non-empty glob pattern.")
    return values


def _inventory_paths(
    root: Path,
    patterns: Sequence[str],
    *,
    recursive: bool,
) -> list[Path]:
    by_relative_path: dict[str, Path] = {}
    for pattern in patterns:
        iterator = root.rglob(pattern) if recursive else root.glob(pattern)
        for path in iterator:
            if not path.is_file():
                continue
            absolute = path.absolute()
            relative = absolute.relative_to(root).as_posix()
            by_relative_path[relative] = absolute
    return [by_relative_path[key] for key in sorted(by_relative_path)]


def _validate_batch_size(value: int) -> int:
    if isinstance(value, bool):
        raise TypeError("batch_size must be a positive integer.")
    try:
        resolved = int(value)
    except (TypeError, ValueError) as exc:
        raise ValueError("batch_size must be a positive integer.") from exc
    if resolved <= 0 or resolved != value:
        raise ValueError("batch_size must be a positive integer.")
    return resolved


def _validate_duckdb_options(
    options: Mapping[str, Any] | None,
) -> dict[str, Any]:
    if options is None:
        return {}
    if not isinstance(options, Mapping):
        raise TypeError("duckdb_options must be a mapping or None.")
    validated: dict[str, Any] = {}
    for raw_name, value in options.items():
        name = str(raw_name)
        if not _OPTION_NAME_RE.fullmatch(name):
            raise ValueError(f"Invalid DuckDB reader option name {name!r}.")
        _duckdb_literal(value)  # validate value shape now
        validated[name] = value
    return validated


def _duckdb_source_sql(
    format: TabularFormat,
    path: Path,
    options: Mapping[str, Any],
) -> str:
    if format == "csv":
        function = "read_csv"
        effective_options = dict(options)
    elif format == "jsonl":
        function = "read_json_auto"
        effective_options = {"format": "newline_delimited", **options}
    elif format == "parquet":
        function = "read_parquet"
        effective_options = dict(options)
    else:  # pragma: no cover - protected by Literal/internal callers
        raise ValueError(f"Unsupported tabular format {format!r}.")

    args = [_duckdb_literal(path.as_posix())]
    args.extend(
        f"{name} = {_duckdb_literal(value)}"
        for name, value in effective_options.items()
    )
    return f"{function}({', '.join(args)})"


def _duckdb_source_sql_many(
    format: TabularFormat,
    paths: Sequence[str],
    options: Mapping[str, Any],
) -> str:
    if not paths:
        raise ValueError("paths must contain at least one source file.")
    if format != "csv":
        raise ValueError("Multi-file source SQL is currently implemented only for CSV.")
    args = [_duckdb_literal(list(paths))]
    args.extend(f"{name} = {_duckdb_literal(value)}" for name, value in options.items())
    return f"read_csv({', '.join(args)})"


def _duckdb_literal(value: Any) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int) and not isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        if not np.isfinite(value):
            raise ValueError("DuckDB reader option floats must be finite.")
        return repr(value)
    if isinstance(value, str):
        return "'" + value.replace("'", "''") + "'"
    if isinstance(value, (list, tuple)):
        return "[" + ", ".join(_duckdb_literal(item) for item in value) + "]"
    if isinstance(value, Mapping):
        items: list[str] = []
        for raw_key, item in value.items():
            key = str(raw_key)
            items.append(f"{_duckdb_literal(key)}: {_duckdb_literal(item)}")
        return "{" + ", ".join(items) + "}"
    raise TypeError(
        "DuckDB reader option values must be JSON-like scalars, lists/tuples, "
        f"or mappings; got {type(value).__name__}."
    )


def _import_duckdb():
    try:
        import duckdb
    except ImportError as exc:  # pragma: no cover - required package in real env
        raise RuntimeError(
            "DuckDB is required for TeAL tabular imports. Install project dependencies."
        ) from exc
    return duckdb


def _write_json(path: Path, data: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(dict(data), indent=2, sort_keys=True)
    temp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=path.parent,
            prefix=f".{path.name}.",
            suffix=".tmp",
            delete=False,
        ) as handle:
            handle.write(payload)
            handle.flush()
            os.fsync(handle.fileno())
            temp_path = Path(handle.name)
        os.replace(temp_path, path)
    finally:
        if temp_path is not None and temp_path.exists():
            temp_path.unlink()
